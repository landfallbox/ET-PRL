#!/usr/bin/env python3
"""
将 build_algorithm_docx.py 生成的算法表格注入 pandoc 转换后的论文 docx。

流程:
    1. 运行 build_algorithm_docx.py → generated/algorithm1_for_merge.docx
  2. pandoc 将 小论文.md → 临时 docx
    3. 为 pandoc 生成的普通论文表格补全所有框线，并合并敏感性表的分组表头
    4. 在临时 docx 中找到算法标题段落 ("Algorithm 1. Event Triggered...")
    5. 删除算法标题至其后第一个 fenced code block 结束之间的旧文本
    6. 将其替换为算法 docx 中的完整表格
    7. 保存到 小论文.docx
"""

from __future__ import annotations

import copy
import argparse
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


BUILD_DIR = Path(__file__).resolve().parent
PAPER_DIR = BUILD_DIR.parent

ALGO_BUILD_SCRIPT = BUILD_DIR / "build_algorithm_docx.py"
ALGO_DOCX = PAPER_DIR / "algorithm1.docx"
GENERATED_DIR = BUILD_DIR / "generated"
MERGE_ALGO_DOCX = GENERATED_DIR / "algorithm1_for_merge.docx"
PAPER_MD = PAPER_DIR / "小论文.md"
OUTPUT_DOCX = PAPER_DIR / "小论文.docx"

ALGO_TITLE_MARKER = "Algorithm 1. Event Triggered Predictive Reinforcement Learning"
ALGO_TITLE_MARKER_EN = "Algorithm 1. Event triggered predictive reinforcement learning"
ALGO_TITLE_MARKER_LEGACY = "Algorithm 1. Event-Triggered Predictive Reinforcement Learning"
ALGO_TITLE_MARKER_EN_LEGACY = "Algorithm 1. Event-triggered predictive reinforcement learning"
NORMAL_TABLE_BORDER_COLOR = "000000"
NORMAL_TABLE_BORDER_SIZE = "4"
SENSITIVITY_PARAMETER_HEADER = "参数值 / Parameter Value"
SENSITIVITY_METRIC_HEADER = "性能指标 / Performance Metrics"

CODE_BLOCK_STYLE_KEYWORDS = (
    "source code",
    "code block",
    "preformatted",
    "verbatim",
)


def _is_heading_paragraph(para) -> bool:
    """判断段落是否是 pandoc 生成的 Word 标题。"""
    return para.style.name.startswith("Heading")


def _is_code_block_paragraph(para) -> bool:
    """判断段落是否来自 Markdown fenced code block。

    pandoc 转换为 docx 后通常不会保留字面量 ``` 标记，而是把 fenced code
    block 转换为 ``Source Code`` 等段落样式。
    """
    style_name = (para.style.name or "").strip().lower()
    return any(keyword in style_name for keyword in CODE_BLOCK_STYLE_KEYWORDS)


def _find_pandoc() -> str:
    """确认 pandoc 在 PATH 上。"""
    import shutil

    pandoc = shutil.which("pandoc")
    if pandoc is None:
        raise RuntimeError(
            "pandoc not found on PATH. Please install pandoc: "
            "https://pandoc.org/installing.html"
        )
    return pandoc


def _get_or_add_child(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def _set_single_border(border_node, *, color: str, size: str) -> None:
    border_node.set(qn("w:val"), "single")
    border_node.set(qn("w:sz"), size)
    border_node.set(qn("w:space"), "0")
    border_node.set(qn("w:color"), color)


def _set_full_grid_borders(table, *, color: str = NORMAL_TABLE_BORDER_COLOR, size: str = NORMAL_TABLE_BORDER_SIZE) -> None:
    tbl_pr = table._tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tbl_pr)

    table_borders = _get_or_add_child(tbl_pr, "w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border_node = _get_or_add_child(table_borders, f"w:{edge}")
        _set_single_border(border_node, color=color, size=size)

    for table_row in table.rows:
        for table_cell in table_row.cells:
            cell_properties = table_cell._tc.get_or_add_tcPr()
            cell_borders = _get_or_add_child(cell_properties, "w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                border_node = _get_or_add_child(cell_borders, f"w:{edge}")
                _set_single_border(border_node, color=color, size=size)


def _set_regular_table_borders(doc: Document) -> int:
    for table in doc.tables:
        _set_full_grid_borders(table)
    return len(doc.tables)


def _cell_text(cell) -> str:
    return " ".join(cell.text.split())


def _set_cell_alignment(cell, alignment) -> None:
    for paragraph in cell.paragraphs:
        paragraph.alignment = alignment


def _table_column_alignments(table) -> list | None:
    if not table.rows:
        return None

    header_texts = [_cell_text(cell) for cell in table.rows[0].cells]
    if not header_texts:
        return None

    first_header = header_texts[0]
    column_count = len(table.columns)

    if first_header in {"策略", "Strategy", "方法", "Method"}:
        return [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * (column_count - 1)

    if first_header == SENSITIVITY_PARAMETER_HEADER:
        return [WD_ALIGN_PARAGRAPH.CENTER] * column_count

    if first_header in {"设备名称", "Equipment"} and column_count >= 6:
        return [
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
        ]

    if first_header in {"模块", "Module"} and column_count >= 4:
        return [
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
        ]

    if first_header == "类别 (Category)" and column_count >= 5:
        return [
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
        ]

    return None


def _apply_regular_table_alignment(doc: Document) -> int:
    aligned_count = 0
    for table in doc.tables:
        alignments = _table_column_alignments(table)
        if alignments is None:
            continue

        for row in table.rows:
            for column_index, cell in enumerate(row.cells):
                if column_index < len(alignments):
                    _set_cell_alignment(cell, alignments[column_index])
        aligned_count += 1

    return aligned_count


def _replace_cell_text(cell, text: str, *, bold: bool = False) -> None:
    cell_properties = cell._tc.tcPr
    for child in list(cell._tc):
        if child is cell_properties:
            continue
        cell._tc.remove(child)

    cell._tc.append(OxmlElement("w:p"))
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold


def _merge_sensitivity_table_group_headers(doc: Document) -> int:
    merged_count = 0
    for table in doc.tables:
        if len(table.rows) < 2 or len(table.columns) < 4:
            continue

        header_cells = table.rows[0].cells
        header_texts = [_cell_text(cell) for cell in header_cells[:4]]
        if header_texts != [
            SENSITIVITY_PARAMETER_HEADER,
            SENSITIVITY_METRIC_HEADER,
            SENSITIVITY_METRIC_HEADER,
            SENSITIVITY_METRIC_HEADER,
        ]:
            continue

        merged_cell = header_cells[1].merge(header_cells[3])
        _replace_cell_text(merged_cell, SENSITIVITY_METRIC_HEADER, bold=True)
        header_cells[0].paragraphs[0].runs[0].bold = True
        merged_count += 1

    return merged_count


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build paper docx from markdown source and inject the formatted algorithm table."
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=PAPER_MD,
        help="Input markdown file. Defaults to docs/papers/小论文.md.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
        help="Output docx path. Defaults to <input_stem>.docx in the same directory.",
    )
    return parser.parse_args()


def main() -> int:
    args = _parse_args()
    input_md = args.input if args.input.is_absolute() else (Path.cwd() / args.input)
    if args.output is not None:
        output_docx = args.output if args.output.is_absolute() else (Path.cwd() / args.output)
    else:
        output_docx = input_md.with_suffix('.docx')

    # --- Step 1: 构建算法 docx ---
    GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    print(f"[1/4] Building algorithm docx...")
    subprocess.run(
        [sys.executable, str(ALGO_BUILD_SCRIPT), "--output", str(MERGE_ALGO_DOCX)],
        check=True,
        cwd=str(BUILD_DIR),
    )

    # --- Step 2: pandoc 将 md 转为临时 docx ---
    pandoc = _find_pandoc()
    with tempfile.TemporaryDirectory() as tmpdir:
        temp_docx = Path(tmpdir) / "paper_temp.docx"

        print(f"[2/4] Converting {input_md.name} to docx via pandoc...")
        subprocess.run(
            [
                pandoc,
                "-f",
                "markdown+tex_math_dollars",
                "-t",
                "docx",
                "-o",
                str(temp_docx),
                str(input_md),
            ],
            check=True,
            cwd=str(PAPER_DIR),
        )

        # --- Step 3: 普通表格补全框线，并合并敏感性表分组表头 ---
        print(f"[3/4] Applying full borders and table header merges...")
        doc = Document(temp_docx)
        bordered_table_count = _set_regular_table_borders(doc)
        print(f"Regular tables updated: {bordered_table_count}")
        merged_header_count = _merge_sensitivity_table_group_headers(doc)
        print(f"Sensitivity table group headers merged: {merged_header_count}")
        aligned_table_count = _apply_regular_table_alignment(doc)
        print(f"Regular tables aligned: {aligned_table_count}")

        # --- Step 4: 注入算法表格 ---
        print(f"[4/4] Injecting algorithm table...")
        algo_doc = Document(MERGE_ALGO_DOCX)

        if not algo_doc.tables:
            print("ERROR: Algorithm docx contains no tables.")
            return 1
        algo_table = algo_doc.tables[0]

        # 定位算法标题段落（大小写不敏感）
        algo_title_para = None
        algo_title_idx = -1
        for i, para in enumerate(doc.paragraphs):
            para_text_lower = para.text.lower()
            if (
                ALGO_TITLE_MARKER.lower() in para_text_lower
                or ALGO_TITLE_MARKER_EN.lower() in para_text_lower
                or ALGO_TITLE_MARKER_LEGACY.lower() in para_text_lower
                or ALGO_TITLE_MARKER_EN_LEGACY.lower() in para_text_lower
            ):
                algo_title_para = para
                algo_title_idx = i
                break

        if algo_title_para is None:
            print(
                f"ERROR: Could not find algorithm title paragraph "
                f"(marker: '{ALGO_TITLE_MARKER}')."
            )
            return 1

        # 查找算法标题之后的第一个 fenced code block。
        # 注意：pandoc 转为 docx 后不会保留字面量 ``` 标记，而是使用
        # Source Code 等段落样式表示代码块。
        code_block_start_idx = -1
        blocking_heading_idx = -1
        for i in range(algo_title_idx + 1, len(doc.paragraphs)):
            para = doc.paragraphs[i]
            if _is_code_block_paragraph(para):
                code_block_start_idx = i
                break
            if _is_heading_paragraph(para):
                blocking_heading_idx = i
                break

        if code_block_start_idx == -1:
            if blocking_heading_idx != -1:
                heading_text = doc.paragraphs[blocking_heading_idx].text
                print(
                    "ERROR: Reached the next heading before finding the "
                    f"Algorithm 1 fenced code block: '{heading_text}'."
                )
            else:
                print(
                    "ERROR: Could not find a fenced code block after the "
                    "Algorithm 1 title. Put Input, Output, and pseudocode "
                    "inside one ```text block."
                )
            return 1

        # 删除范围截止到第一个连续代码块区域结束。
        algo_end_idx = code_block_start_idx + 1
        while (
            algo_end_idx < len(doc.paragraphs)
            and _is_code_block_paragraph(doc.paragraphs[algo_end_idx])
        ):
            algo_end_idx += 1

        # 算法标题的前一个元素 → 将表格插入到它之后
        prev_element = algo_title_para._element.getprevious()
        if prev_element is None:
            print("ERROR: Algorithm title is the first element in the document body.")
            return 1

        # 收集待删除的段落元素（算法标题 + 标题后第一个代码块结束前的内容）
        elements_to_remove = [
            doc.paragraphs[i]._element
            for i in range(algo_title_idx, algo_end_idx)
        ]

        # 深拷贝算法表格并插入
        tbl_copy = copy.deepcopy(algo_table._tbl)
        prev_element.addnext(tbl_copy)

        # 删除旧的算法文本段落
        for elem in elements_to_remove:
            elem.getparent().remove(elem)

        # 保存最终文档
        doc.save(output_docx)

    print(f"Done! Created {output_docx}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
