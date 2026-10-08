"""Validate the merged docx output."""
from __future__ import annotations

from pathlib import Path
from docx import Document

DOCX_PATH = Path("docs/papers/et-prl/小论文.docx")
doc = Document(DOCX_PATH)

print(f"Total tables: {len(doc.tables)}")
print(f"Total paragraphs: {len(doc.paragraphs)}")

# Check "summarized below" exists
found_summarized = False
for i, para in enumerate(doc.paragraphs):
    if "summarized below" in para.text:
        found_summarized = True
        print(f'  [OK] Found "summarized below" at paragraph {i}')
        break
if not found_summarized:
    print("  [WARN] 'summarized below' not found (might be on next page)")

# Check no old Algorithm 1 markdown/text block remains in paragraphs
old_algorithm_markers = (
    "Algorithm 1. Event-Triggered",
    "Input: Trained DQN policy",
    "Output: Control setpoint sequence",
    "Initialization:",
    "for t = 1, ..., T do",
)
remaining_old_paras = [
    p.text
    for p in doc.paragraphs
    if any(marker in p.text for marker in old_algorithm_markers)
]
if not remaining_old_paras:
    print("  [OK] No old Algorithm 1 text/code paragraphs remain")
else:
    print(f"  [WARN] Found {len(remaining_old_paras)} old Algorithm 1 paragraphs")
    for text in remaining_old_paras[:3]:
        print(f'       - "{text[:100]}..."')

# Check table content
if doc.tables:
    t = doc.tables[0]
    print(f"  Algorithm table: {len(t.rows)} rows x {len(t.columns)} cols")
    first_cell_text = t.rows[0].cells[0].text
    if "Algorithm 1" in first_cell_text:
        print(f'  [OK] Table title: "{first_cell_text[:80]}..."')
    else:
        print(f'  [INFO] Table first cell: "{first_cell_text[:80]}..."')
else:
    print("  [FAIL] No tables found in output docx!")

# Check Section 4 heading exists (document structure intact)
has_section_4 = any("4." in p.text and "案例系统" in p.text for p in doc.paragraphs)
print(f"  {'[OK]' if has_section_4 else '[FAIL]'} Section 4 heading: {'found' if has_section_4 else 'missing'}")

print("\nValidation complete.")
