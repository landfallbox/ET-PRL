from __future__ import annotations

import argparse
import copy
import subprocess
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree


PANDOC = "pandoc"

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
}

PAGE_WIDTH_DXA = 9360
COL_WIDTHS_DXA = [360, 360, 360, 8280]
SERIF_FONT = "Times New Roman"
BORDER_COLOR = "7F7F7F"
BORDER_SIZE = "6"


Segment = tuple[str, str]


FORMULAS = [
    r"\pi_{\mathrm{DQN}}",
    r"\Omega=\{w_{\mathrm{s}},w_{\mathrm{m}},w_{\mathrm{l}},q,\kappa,\omega_q,\lambda_{\mathrm{local}},\lambda_{\mathrm{global}},\alpha_{\mathrm{local}},b_{\mathrm{bias}},W,N_{\mathrm{ref}},\nu,m_{\mathrm{hys}}\}",
    r"\Delta_{\min}",
    r"T",
    r"\{a_t\}_{t=1}^{T}",
    r"a_0",
    r"t_{\mathrm{last}}=-\infty",
    r"A_t^{(W)}",
    r"t=1,\ldots,T",
    r"Q_{\mathrm{load}}^t",
    r"T_{\mathrm{wb}}^t",
    r"\hat{Q}_{\mathrm{load}}^{t+1}",
    r"s_t=[Q_{\mathrm{load}}^t,T_{\mathrm{wb}}^t,\hat{Q}_{\mathrm{load}}^{t+1}]",
    r"A_{\mathrm{short}}(s_t)",
    r"A_{\mathrm{medium}}(s_t)",
    r"A_{\mathrm{long}}(s_t)",
    r"A(s_t)",
    r"\tau_t^*",
    r"A(s_t)>\tau_t^*",
    r"t-t_{\mathrm{last}}>\Delta_{\min}",
    r"a_t=\pi_{\mathrm{DQN}}(s_t)",
    r"t_{\mathrm{last}}=t",
    r"a_t=a_{t-1}",
    r"a_t",
    r"A(s_t)",
    r"A_t^{(W)}",
]


ROWS: list[dict[str, object]] = [
    {
        "kind": "title",
        "indent": 0,
        "segments": [("text", "Algorithm 1. Event Triggered Predictive Reinforcement Learning with Unsupervised Dynamic Event Gating (ET-PRL) Online Control Procedure")],
    },
    {
        "kind": "meta",
        "indent": 0,
        "segments": [
            ("bold", "Input: "),
            ("text", "Trained DQN policy "),
            ("math", FORMULAS[0]),
            ("text", "; gate parameters "),
            ("math", FORMULAS[1]),
            ("text", " defined in Eqs. (7)-(16); minimum trigger interval "),
            ("math", FORMULAS[2]),
            ("text", "; total control horizon "),
            ("math", FORMULAS[3]),
            ("text", "."),
        ],
    },
    {
        "kind": "meta",
        "indent": 0,
        "segments": [
            ("bold", "Output: "),
            ("text", "Control setpoint sequence "),
            ("math", FORMULAS[4]),
            ("text", "."),
        ],
    },
    {"kind": "code", "indent": 0, "segments": [("text", "Initialization:")]},
    {"kind": "code", "indent": 1, "segments": [("text", "Set the initial control action "), ("math", FORMULAS[5])]},
    {"kind": "code", "indent": 1, "segments": [("text", "Set the last trigger time "), ("math", FORMULAS[6])]},
    {
        "kind": "code",
        "indent": 1,
        "segments": [("text", "Initialize the anomaly score sliding window "), ("math", FORMULAS[7])],
    },
    {"kind": "code", "indent": 0, "segments": [("text", "for "), ("math", FORMULAS[8]), ("text", " do")]},
    {
        "kind": "code",
        "indent": 1,
        "segments": [
            ("text", "Collect "),
            ("math", FORMULAS[9]),
            ("text", ", "),
            ("math", FORMULAS[10]),
            ("text", ", and "),
            ("math", FORMULAS[11]),
        ],
    },
    {
        "kind": "code",
        "indent": 1,
        "segments": [("text", "Construct the predictive state "), ("math", FORMULAS[12])],
    },
    {
        "kind": "code",
        "indent": 1,
        "segments": [
            ("text", "Compute the scale-specific anomaly scores "),
            ("math", FORMULAS[13]),
            ("text", ", "),
            ("math", FORMULAS[14]),
            ("text", ", and "),
            ("math", FORMULAS[15]),
        ],
    },
    {
        "kind": "code",
        "indent": 1,
        "segments": [
            ("text", "Compute the multi-scale fused anomaly score "),
            ("math", FORMULAS[16]),
            ("text", " according to Eq. (8)"),
        ],
    },
    {
        "kind": "code",
        "indent": 1,
        "segments": [("text", "Compute the candidate threshold, local threshold, and global threshold according to Eqs. (11)-(14)")],
    },
    {
        "kind": "code",
        "indent": 1,
        "segments": [
            ("text", "Compute the adaptive triggering threshold "),
            ("math", FORMULAS[17]),
            ("text", " according to Eq. (15)"),
        ],
    },
    {
        "kind": "code",
        "indent": 1,
        "segments": [
            ("text", "if "),
            ("math", FORMULAS[18]),
            ("text", " and "),
            ("math", FORMULAS[19]),
            ("text", " then"),
        ],
    },
    {
        "kind": "code",
        "indent": 2,
        "segments": [
            ("text", "Infer a new control setpoint using the trained DQN policy: "),
            ("math", FORMULAS[20]),
        ],
    },
    {
        "kind": "code",
        "indent": 2,
        "segments": [("text", "Update the last trigger time: "), ("math", FORMULAS[21])],
    },
    {"kind": "code", "indent": 1, "segments": [("text", "else")]},
    {
        "kind": "code",
        "indent": 2,
        "segments": [("text", "Hold the previous control setpoint: "), ("math", FORMULAS[22])],
    },
    {"kind": "code", "indent": 1, "segments": [("text", "end if")]},
    {"kind": "code", "indent": 1, "segments": [("text", "Execute "), ("math", FORMULAS[23])]},
    {
        "kind": "code",
        "indent": 1,
        "segments": [
            ("text", "Append "),
            ("math", FORMULAS[24]),
            ("text", " to "),
            ("math", FORMULAS[25]),
            ("text", " for use in the next threshold update"),
        ],
    },
    {"kind": "code", "indent": 0, "segments": [("text", "end for")]},
]


def run_pandoc_for_formulas(work_dir: Path) -> list[etree._Element]:
    work_dir.mkdir(parents=True, exist_ok=True)
    math_markdown = work_dir / "formula_source.md"
    math_docx = work_dir / "formula_source.docx"
    math_markdown.write_text("\n\n".join(f"${formula}$" for formula in FORMULAS), encoding="utf-8")
    subprocess.run(
        [PANDOC, "-f", "markdown+tex_math_dollars", "-t", "docx", "--standalone", str(math_markdown), "-o", str(math_docx)],
        check=True,
    )
    with zipfile.ZipFile(math_docx) as archive:
        xml = archive.read("word/document.xml")
    root = etree.fromstring(xml)
    equations = root.xpath(".//m:oMath", namespaces=NS)
    if len(equations) != len(FORMULAS):
        raise RuntimeError(f"Expected {len(FORMULAS)} OMML equations, found {len(equations)}")
    return [copy.deepcopy(eq) for eq in equations]


def set_cell_margins(cell, top=30, start=120, bottom=30, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")
        node.set(qn("w:sz"), "0")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "auto")


def set_cell_border(cell, edge: str, *, color: str = BORDER_COLOR, size: str = BORDER_SIZE) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    node = borders.find(qn(f"w:{edge}"))
    if node is None:
        node = OxmlElement(f"w:{edge}")
        borders.append(node)
    node.set(qn("w:val"), "single")
    node.set(qn("w:sz"), size)
    node.set(qn("w:space"), "0")
    node.set(qn("w:color"), color)


def set_algorithm_rules(table) -> None:
    for cell in table.rows[0].cells:
        set_cell_border(cell, "top")
        set_cell_border(cell, "bottom")

    for cell in table.rows[-1].cells:
        set_cell_border(cell, "bottom")


def set_table_geometry(table) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(PAGE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tbl_grid = table._tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in COL_WIDTHS_DXA:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(COL_WIDTHS_DXA[idx]))
            tc_w.set(qn("w:type"), "dxa")


def set_document_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = SERIF_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), SERIF_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), SERIF_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), SERIF_FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE

    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def add_segments(paragraph, segments: list[Segment], equations: dict[str, etree._Element], kind: str) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    for segment_kind, value in segments:
        if segment_kind == "math":
            paragraph._p.append(copy.deepcopy(equations[value]))
            continue

        run = paragraph.add_run(value)
        run.font.name = SERIF_FONT
        run._element.rPr.rFonts.set(qn("w:ascii"), SERIF_FONT)
        run._element.rPr.rFonts.set(qn("w:hAnsi"), SERIF_FONT)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), SERIF_FONT)
        run.font.size = Pt(10.5 if kind != "title" else 11)
        run.font.color.rgb = RGBColor(0, 0, 0)
        if segment_kind == "bold":
            run.font.bold = True


def build_document(equations: dict[str, etree._Element], output: Path) -> None:
    document = Document()
    set_document_styles(document)

    table = document.add_table(rows=len(ROWS), cols=4)
    table.autofit = False
    set_table_geometry(table)
    set_table_borders(table)

    for row_idx, row_info in enumerate(ROWS):
        row = table.rows[row_idx]
        kind = str(row_info["kind"])
        indent = int(row_info["indent"])
        segments = row_info["segments"]

        for cell_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)

        target_cell = row.cells[indent]
        if indent < 3:
            target_cell = target_cell.merge(row.cells[3])
        elif kind in {"title", "meta"}:
            target_cell = row.cells[0].merge(row.cells[3])

        if kind in {"title", "meta"}:
            target_cell = row.cells[0]

        paragraph = target_cell.paragraphs[0]
        add_segments(paragraph, segments, equations, kind)

    set_algorithm_rules(table)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def audit_docx(output: Path) -> tuple[int, int]:
    with zipfile.ZipFile(output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    return xml.count("<m:oMath"), xml.count("<w:tr")


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build the Algorithm 1 Word table.")
    parser.add_argument(
        "--output",
        type=Path,
        required=True,
        help="Output .docx path.",
    )
    parser.add_argument(
        "--work-dir",
        type=Path,
        required=True,
        help="Directory for temporary Pandoc formula files.",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> None:
    args = parse_args(argv)
    output = args.output.resolve()
    formula_nodes = run_pandoc_for_formulas(args.work_dir.resolve())
    equations = dict(zip(FORMULAS, formula_nodes))
    build_document(equations, output)
    math_count, row_count = audit_docx(output)
    print(f"Created {output}")
    print(f"Rows: {row_count}")
    print(f"OMML equation elements: {math_count}")


if __name__ == "__main__":
    sys.exit(main())
